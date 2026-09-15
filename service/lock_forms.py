#
# フォームのセル保護
#
import openpyxl
from openpyxl.styles import Protection
#
from config import IO_DATA


def フォームをロックする():
    """
    商品マスタ・仕入れフォームの、キー列に自動採番してロックし、
    販売登録フォームの、数式が入った列をロックする
    """
    wb = openpyxl.load_workbook(IO_DATA['商品別売上'])

    _商品マスタをロックする(wb['商品マスタ'])
    _仕入れフォームをロックする(wb['仕入れフォーム'])
    _販売登録フォームをロックする(wb['販売登録フォーム'])

    wb.save(IO_DATA['商品別売上'])


#-----
# プライベート関数
#-----

def _商品マスタをロックする(ws):
    """
    商品キーID(B列)が空欄の行に自動採番し、
    採番済みのセルをロックする
    """
    最大ID = 0
    for r in range(2, ws.max_row + 1):
        値 = ws.cell(row=r, column=2).value
        if isinstance(値, (int, float)):
            最大ID = max(最大ID, int(値))

    for r in range(2, ws.max_row + 1):
        商品名 = ws.cell(row=r, column=3).value
        IDセル = ws.cell(row=r, column=2)

        if 商品名 is not None and IDセル.value is None:
            最大ID += 1
            IDセル.value = 最大ID

        # 採番済みなら、ロックする
        IDセル.protection = Protection(locked=(IDセル.value is not None))

        # 商品名・仕入れ先・品目は、手入力できるままにする
        for col in (3, 4, 5):
            ws.cell(row=r, column=col).protection = Protection(locked=False)

    ws.protection.sheet = True


def _仕入れフォームをロックする(ws):
    """
    仕入れキー(A列)が空欄の行に自動採番し、
    採番済みのセルをロックする
    """
    最大キー = 0
    for r in range(2, ws.max_row + 1):
        値 = ws.cell(row=r, column=1).value
        if isinstance(値, (int, float)):
            最大キー = max(最大キー, int(値))

    for r in range(2, ws.max_row + 1):
        商品キーID = ws.cell(row=r, column=2).value
        キーセル = ws.cell(row=r, column=1)

        if 商品キーID is not None and キーセル.value is None:
            最大キー += 1
            キーセル.value = 最大キー

        キーセル.protection = Protection(locked=(キーセル.value is not None))

        # 商品キーID・入荷日・入荷数量・仕入れ単価・定価は、手入力できるままにする
        for col in (2, 3, 4, 5, 6):
            ws.cell(row=r, column=col).protection = Protection(locked=False)

    ws.protection.sheet = True


def _販売登録フォームをロックする(ws):
    """
    数式が入っている列(商品キーID・ブルワリー・商品名・品目)を、
    3行目以降だけロックする。
    1行目(出荷先)・2行目(見出し・日付)は、入力できるように、ロックしない。
    """
    # 1行目・2行目は、全列、ロックしない(出荷先・日付を、入力できるように)
    for r in (1, 2):
        for col in range(1, ws.max_column + 1):
            ws.cell(row=r, column=col).protection = Protection(locked=False)

    # 3行目以降だけ、A・C・D・E列(数式の列)を、ロックする
    for r in range(3, ws.max_row + 1):
        for col in (1, 3, 4, 5):  # A, C, D, E
            セル = ws.cell(row=r, column=col)
            セル.protection = Protection(locked=(セル.value is not None))

        # 欠品・期首在庫・日付列(出荷数量)は、手入力できるままにする
        for col in [2, 6] + list(range(7, ws.max_column + 1)):
            ws.cell(row=r, column=col).protection = Protection(locked=False)

    ws.protection.sheet = True